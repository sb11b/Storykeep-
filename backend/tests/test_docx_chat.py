from __future__ import annotations

import io
import unittest
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile
from types import SimpleNamespace

from app.services.docx_chat import (
    DOCX_READ_ERROR,
    DOC_HINT,
    MACRO_HINT,
    extract_chat_docx,
    inspect_docx_bytes,
    sanitize_docx_html,
)
from app.services.chat_attachments import extract_text_for_media, merge_attachment_text


def _sample_docx(*paragraphs: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(paragraphs[0], level=1)
    for line in paragraphs[1:]:
        document.add_paragraph(line)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _zip_with(names: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        for name, payload in names.items():
            archive.writestr(name, payload)
    return buf.getvalue()


class DocxChatTests(unittest.TestCase):
    def test_mammoth_extracts_headings_and_paragraphs(self):
        payload = _sample_docx("Helium ash", "The reactor is online.", "Keep the column.")
        text, html = extract_chat_docx(payload, "notes.docx")
        self.assertIn("Helium ash", text)
        self.assertIn("The reactor is online.", text)
        self.assertIn("<h1>", html)
        self.assertIn("<p>", html)
        self.assertNotIn("<script", html)

    def test_sanitizer_drops_scripts_iframes_and_non_https_links(self):
        cleaned = sanitize_docx_html(
            '<p>Body</p><script>alert(1)</script><iframe src="https://evil.example"></iframe>'
            '<p><a href="javascript:alert(1)">x</a><a href="https://example.com/ok">ok</a></p>'
        )
        self.assertIn("Body", cleaned)
        self.assertNotIn("script", cleaned.lower())
        self.assertNotIn("iframe", cleaned.lower())
        self.assertNotIn("javascript:", cleaned.lower())
        self.assertIn("https://example.com/ok", cleaned)

    def test_rejects_doc_ole_and_macros(self):
        with self.assertRaises(ValueError) as doc_err:
            inspect_docx_bytes(b"unused", "old.doc")
        self.assertEqual(str(doc_err.exception), DOC_HINT)
        with self.assertRaises(ValueError) as ole_err:
            inspect_docx_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 20, "macro.docx")
        self.assertEqual(str(ole_err.exception), MACRO_HINT)
        poisoned = _zip_with(
            {
                "word/document.xml": b"<w:document/>",
                "word/vbaProject.bin": b"MZ",
            }
        )
        with self.assertRaises(ValueError) as macro_err:
            inspect_docx_bytes(poisoned, "macro.docx")
        self.assertEqual(str(macro_err.exception), MACRO_HINT)
        embedded = _zip_with(
            {
                "word/document.xml": b"<w:document/>",
                "word/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0",
            }
        )
        with self.assertRaises(ValueError):
            inspect_docx_bytes(embedded, "ole.docx")

    def test_extract_text_for_media_puts_docx_on_the_turn(self):
        payload = _sample_docx("DAT-200 homework", "Gravity is a field.")
        with NamedTemporaryFile(suffix=".docx", delete=False) as handle:
            handle.write(payload)
            path = Path(handle.name)
        try:
            row = SimpleNamespace(
                filename="homework.docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                storage_path=str(path),
            )
            text = extract_text_for_media(row)
            self.assertIn("DAT-200 homework", text)
            self.assertIn("Gravity is a field.", text)
            merged = merge_attachment_text("Summarize this", [{"filename": "homework.docx", "kind": "file", "extract_text": text}], include_extracts=True)
            self.assertIn("Gravity is a field.", merged)
        finally:
            path.unlink(missing_ok=True)

    def test_unreadable_docx_returns_chat_error(self):
        with NamedTemporaryFile(suffix=".docx", delete=False) as handle:
            handle.write(b"PK not a document")
            path = Path(handle.name)
        try:
            row = SimpleNamespace(filename="broken.docx", content_type="application/octet-stream", storage_path=str(path))
            self.assertEqual(extract_text_for_media(row), DOCX_READ_ERROR)
        finally:
            path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
