from __future__ import annotations

import io
import json
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.services.vault_paths import windows_safe_component

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DEFAULT_FILENAME = "junior-note.docx"
EMPTY_WORD_BODY = "That reply has no text to put in Word."
BUILD_WORD_FAIL = "Couldn't build Word"
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_UL = re.compile(r"^[-*+]\s+(.*)$")
_OL = re.compile(r"^\d+[.)]\s+(.*)$")
_FENCE = re.compile(r"^```")
_FENCE_BLOCK = re.compile(r"```[\s\S]*?```")
_IMAGE = re.compile(r"!\[[^\]]*\]\([^)]*\)")
_LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
_CODE = re.compile(r"`([^`]+)`")
_REF_HEAD = re.compile(r"^(references?|works cited|bibliography)\b", re.I)
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


_KEEP_NOTES_LINE = re.compile(
    r"^(?:[-*•]\s+)?(?:please\s+)?"
    r"(?:"
    r"use add to notes if you want this kept"
    r"|if you want this kept,?\s+use add to notes"
    r"|if you(?:'d| would)? like this kept,?\s+use add to notes"
    r"|if you want to keep this(?: reply)?,?\s+use add to notes"
    r"|add to notes if you want this(?: reply)? kept"
    r"|use add to notes to keep this(?: reply)?"
    r")\.?$",
    re.I,
)
_KEEP_NOTES_TAIL = re.compile(
    r"(?:\s+)"
    r"(?:use add to notes if you want this kept"
    r"|if you want this kept,?\s+use add to notes"
    r"|if you(?:'d| would)? like this kept,?\s+use add to notes"
    r"|add to notes if you want this(?: reply)? kept)\.?\s*$",
    re.I,
)


def strip_keep_notes_cta(content: str) -> str:
    """Drop system keep/notes footers so Copy/Word/clipboard stay the answer body."""
    text = (content or "").replace("\r\n", "\n")
    lines = [line for line in text.split("\n") if not _KEEP_NOTES_LINE.match(line.strip())]
    cleaned = "\n".join(lines)
    cleaned = _KEEP_NOTES_TAIL.sub("", cleaned)
    return re.sub(r"\n{3,}", "\n\n", cleaned).strip()


def visible_reply_text(content: str) -> str:
    text = strip_keep_notes_cta(content)
    if not text:
        return ""
    text = _IMAGE.sub("", text)
    text = _LINK.sub(r"\1", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def prose_for_word(content: str) -> str:
    """Visible reply with fenced tool/code blocks removed."""
    text = visible_reply_text(content)
    if not text:
        return ""
    text = _FENCE_BLOCK.sub("", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def has_word_body(content: str) -> bool:
    text = prose_for_word(content)
    if not text:
        return False
    compact = re.sub(r"\s+", " ", text).strip()
    if not compact:
        return False
    if compact[0] in "{[" and compact[-1] in "}]":
        try:
            json.loads(compact)
            return False
        except (ValueError, TypeError):
            pass
    return True


def first_heading(content: str) -> str | None:
    for raw in visible_reply_text(content).split("\n"):
        heading = _HEADING.match(raw.strip())
        if not heading:
            continue
        line = _CODE.sub(r"\1", _BOLD.sub(r"\1", _ITALIC.sub(r"\1", heading.group(2).strip())))
        line = " ".join(line.split())
        if line:
            return line[:120]
    return None


def docx_title(content: str) -> str:
    heading = first_heading(content)
    if heading:
        return heading
    for raw in visible_reply_text(content).split("\n"):
        line = raw.strip()
        if not line or _FENCE.match(line):
            continue
        line = _CODE.sub(r"\1", _BOLD.sub(r"\1", _ITALIC.sub(r"\1", line)))
        line = " ".join(line.split())
        if line:
            return line[:120]
    return "Junior reply"


def text_filename(content: str, ext: str) -> str:
    suffix = (ext or "txt").lstrip(".").lower()
    if suffix not in {"md", "txt"}:
        suffix = "txt"
    heading = first_heading(content)
    stem = windows_safe_component(heading or "") if heading else "junior-note"
    stem = stem or "junior-note"
    return f"{stem}.{suffix}"


def docx_filename(content: str) -> str:
    stem = windows_safe_component(docx_title(content))
    if not stem or stem == "_":
        stem = "junior-note"
    if stem.lower().endswith(".docx"):
        return stem
    return f"{stem}.docx"


def _set_run_font(run, *, bold: bool = False, italic: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = run._element.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:cs"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "Times New Roman")


def _add_styled_text(paragraph, text: str, *, heading: bool = False) -> None:
    remaining = text or ""
    index = 0
    for match in _INLINE.finditer(remaining):
        if match.start() > index:
            run = paragraph.add_run(remaining[index : match.start()])
            _set_run_font(run, bold=heading)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, bold=heading)
        else:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, italic=True)
        index = match.end()
    if index < len(remaining):
        run = paragraph.add_run(remaining[index:])
        _set_run_font(run, bold=heading)
    if not paragraph.runs:
        run = paragraph.add_run("")
        _set_run_font(run, bold=heading)
    if heading:
        for run in paragraph.runs:
            run.bold = True


def _apply_page(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)


def _hanging_indent(paragraph) -> None:
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = paragraph._p.makeelement(qn("w:ind"), {})
        p_pr.append(ind)
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "720")


def build_message_docx(content: str) -> bytes:
    text = visible_reply_text(content)
    if not text or not has_word_body(content):
        raise ValueError(EMPTY_WORD_BODY)
    document = Document()
    _apply_page(document)
    in_references = False
    in_fence = False
    fence_lines: list[str] = []
    lines = text.split("\n")
    index = 0
    while index < len(lines):
        line = lines[index]
        if _FENCE.match(line.strip()):
            if in_fence:
                paragraph = document.add_paragraph()
                _add_styled_text(paragraph, "\n".join(fence_lines))
                fence_lines = []
                in_fence = False
            else:
                in_fence = True
            index += 1
            continue
        if in_fence:
            fence_lines.append(line)
            index += 1
            continue
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        heading = _HEADING.match(stripped)
        if heading:
            title = heading.group(2).strip()
            in_references = bool(_REF_HEAD.match(_CODE.sub(r"\1", _BOLD.sub(r"\1", title))))
            paragraph = document.add_paragraph()
            _add_styled_text(paragraph, title, heading=True)
            index += 1
            continue
        if _REF_HEAD.match(stripped) and len(stripped) < 40:
            in_references = True
            paragraph = document.add_paragraph()
            _add_styled_text(paragraph, stripped, heading=True)
            index += 1
            continue
        bullet = _UL.match(stripped)
        numbered = _OL.match(stripped)
        paragraph = document.add_paragraph()
        if in_references and not bullet and not numbered:
            _hanging_indent(paragraph)
        if bullet:
            _add_styled_text(paragraph, f"• {bullet.group(1).strip()}")
        elif numbered:
            _add_styled_text(paragraph, stripped)
        else:
            _add_styled_text(paragraph, stripped)
        index += 1
    if in_fence and fence_lines:
        paragraph = document.add_paragraph()
        _add_styled_text(paragraph, "\n".join(fence_lines))
    payload = io.BytesIO()
    document.save(payload)
    data = payload.getvalue()
    if not data.startswith(b"PK"):
        raise ValueError(BUILD_WORD_FAIL)
    with ZipFile(io.BytesIO(data)) as archive:
        if "word/document.xml" not in archive.namelist():
            raise ValueError(BUILD_WORD_FAIL)
    return data


def document_xml(payload: bytes) -> str:
    with ZipFile(io.BytesIO(payload)) as archive:
        return archive.read("word/document.xml").decode("utf-8")
