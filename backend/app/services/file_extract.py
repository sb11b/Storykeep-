from __future__ import annotations

import csv
import html as html_lib
import io
import re
import zipfile
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET


ALLOWED_SUFFIXES = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".csv",
    ".rtf",
    ".odt",
    ".epub",
}

DOC_HINT = "Legacy Word .doc is not supported. Save as PDF or .docx and upload again."


def suffix_of(filename: str) -> str:
    return PurePosixPath((filename or "").replace("\\", "/")).suffix.lower()


def extract_document(filename: str, payload: bytes, *, max_pdf_pages: int | None = None) -> tuple[str, str]:
    name = PurePosixPath((filename or "upload").replace("\\", "/")).name or "upload"
    suffix = suffix_of(name)
    if suffix == ".doc":
        raise ValueError(DOC_HINT)
    if suffix == ".zip":
        raise ValueError("Obsidian vault zips go in Collect → Vault. This File tab is one PDF, Word, PowerPoint, or text file.")
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError(
            f"Cannot extract {suffix or 'that file'}. Upload PDF, Word (.docx), PowerPoint (.pptx), text, markdown, HTML, CSV, RTF, ODT, or EPUB."
        )
    text = ""
    try:
        if suffix == ".pdf":
            text = _pdf(payload, max_pages=max_pdf_pages)
        elif suffix == ".docx":
            text = _docx(payload)
        elif suffix == ".pptx":
            text = _pptx(payload)
        elif suffix in {".txt", ".md", ".markdown"}:
            text = payload.decode("utf-8", errors="replace")
        elif suffix in {".html", ".htm"}:
            text = _html(payload)
        elif suffix == ".csv":
            text = _csv(payload)
        elif suffix == ".rtf":
            text = _rtf(payload)
        elif suffix == ".odt":
            text = _odt(payload)
        elif suffix == ".epub":
            text = _epub(payload)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not read that {suffix} file.") from exc
    text = _clean(text)
    title = PurePosixPath(name).stem.replace("_", " ").strip() or name
    return title, text


def _clean(text: str) -> str:
    text = (text or "").replace("\x00", "")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(payload: bytes, max_pages: int | None = None) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(payload))
    total = len(reader.pages)
    pages: list[str] = []
    stopped = False
    for index, page in enumerate(reader.pages):
        if max_pages is not None and index >= max_pages:
            stopped = True
            break
        body = page.extract_text() or ""
        pages.append(f"--- page {index + 1} of {total} ---\n{body}".rstrip())
    text = "\n\n".join(pages)
    if stopped and max_pages is not None:
        text += f"\n\n[Extract stopped after {max_pages} pages of {total}.]"
    return text


def _docx(payload: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(payload))
    blocks = [para.text for para in document.paragraphs if para.text and para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _pptx(payload: bytes) -> str:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(payload))
    slides = []
    for index, slide in enumerate(deck.slides, start=1):
        bits = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = shape.text_frame.text.strip()
                if value:
                    bits.append(value)
        if bits:
            slides.append(f"## Slide {index}\n\n" + "\n\n".join(bits))
    return "\n\n".join(slides)


def _html(payload: bytes) -> str:
    raw = payload.decode("utf-8", errors="replace")
    try:
        from app.services.extractor import extract_html

        _, text = extract_html(raw)
        if text and text.strip():
            return text
    except Exception:
        pass
    plain = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
    plain = re.sub(r"(?s)<[^>]+>", " ", plain)
    return html_lib.unescape(plain)


def _csv(payload: bytes) -> str:
    text = payload.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = []
    for row in reader:
        cells = [cell.strip() for cell in row if cell.strip()]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _rtf(payload: bytes) -> str:
    raw = payload.decode("latin-1", errors="replace")
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d*[ ]?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    return re.sub(r"\s+", " ", raw).strip()


def _odt(payload: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        xml = archive.read("content.xml")
    root = ET.fromstring(xml)
    texts = [node.text for node in root.iter() if node.text and node.text.strip()]
    return "\n".join(texts)


def _epub(payload: bytes) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = [name for name in archive.namelist() if name.lower().endswith((".xhtml", ".html", ".htm"))]
        for name in sorted(names):
            chunks.append(_html(archive.read(name)))
    return "\n\n".join(chunks)
